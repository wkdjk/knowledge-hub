#!/usr/bin/env python3
"""Convert the QIA quarantine statistics workbook into the site's CSV files.

Usage:
    python3 scripts/convert_qia.py "sources/Korean import statistics for deer velvet.xlsx"

Outputs (always written relative to the repo root):
    data/imports-annual.csv  - source of record: year x country x form (dried/frozen), kg
    data/imports.csv         - chart file: year x country, dried-equivalent tonnes
                               (dried kg + 0.33 x frozen kg, complete years only)

The same input always produces byte-identical output (fixed sort order, no
timestamps). After running, a summary is printed for eyeball verification.
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
SHEET_NAME = "Quarantine stats"
SECTIONS = {"Dried velvet (kg)": "dried", "Frozen velvet (kg)": "frozen"}
FROZEN_TO_DRIED = 0.33  # dried-equivalent convention used in the workbook itself


def fail(msg):
    sys.exit(f"ERROR: {msg}")


def parse_year(header):
    """Return (year, is_partial) from a header like 2013 or '2026 (~Jun)'."""
    if isinstance(header, (int, float)):
        return int(header), False
    m = re.match(r"^\s*(\d{4})\s*(\(.*\))?\s*$", str(header))
    if not m:
        return None, None
    return int(m.group(1)), m.group(2) is not None


def main():
    if len(sys.argv) != 2:
        fail(f"usage: {sys.argv[0]} <QIA workbook.xlsx>")
    src = Path(sys.argv[1])
    if not src.exists():
        fail(f"input file not found: {src}")

    wb = openpyxl.load_workbook(src, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        fail(f"expected sheet '{SHEET_NAME}' not found; sheets are {wb.sheetnames}")
    ws = wb[SHEET_NAME]

    # Walk column B: section titles, then a Country header row, then country rows
    # ending at 'Total'. Year columns are read from each section's header row.
    records = []          # (year, country, form, kg, partial)
    section_form = None
    year_cols = {}        # col index -> (year, partial)
    in_table = False
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 2).value
        label = str(label).strip() if label is not None else ""
        if label in SECTIONS:
            section_form = SECTIONS[label]
            in_table = False
            continue
        if section_form and label == "Country":
            year_cols = {}
            for c in range(3, ws.max_column + 1):
                year, partial = parse_year(ws.cell(r, c).value)
                if year and c < 17:  # columns beyond ~17 are the rolling-year block
                    year_cols[c] = (year, partial)
            if not year_cols:
                fail(f"no year columns found in header at row {r}")
            in_table = True
            continue
        if in_table:
            if label == "" or label == "Total":
                if label == "Total":
                    in_table = False
                    section_form = None
                continue
            for c, (year, partial) in year_cols.items():
                v = ws.cell(r, c).value
                if v is None or v == "":
                    continue
                try:
                    kg = float(v)
                except (TypeError, ValueError):
                    fail(f"non-numeric value {v!r} at row {r} col {c}")
                records.append((year, label, section_form, kg, partial))

    if not records:
        fail("no data rows parsed - has the workbook layout changed?")

    records.sort(key=lambda x: (x[0], x[1], x[2]))

    annual_path = REPO_ROOT / "data" / "imports-annual.csv"
    with open(annual_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "country", "form", "kg", "coverage"])
        for year, country, form, kg, partial in records:
            w.writerow([year, country, form, f"{kg:g}",
                        "partial" if partial else "full"])

    # Chart file: dried-equivalent tonnes per year x country, complete years only.
    eq = {}
    partial_years = sorted({y for y, *_ , p in records if p})
    for year, country, form, kg, partial in records:
        if partial:
            continue
        factor = 1.0 if form == "dried" else FROZEN_TO_DRIED
        eq[(year, country)] = eq.get((year, country), 0.0) + kg * factor

    imports_path = REPO_ROOT / "data" / "imports.csv"
    with open(imports_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["year", "country", "tonnes"])
        for (year, country) in sorted(eq):
            w.writerow([year, country, f"{eq[(year, country)] / 1000:.1f}"])

    # ---- verification summary ----
    years = sorted({y for y, *_ in records})
    print(f"input : {src}")
    print(f"parsed: {len(records)} data cells "
          f"({len({c for _, c, *_ in records})} countries, "
          f"years {years[0]}-{years[-1]})")
    if partial_years:
        print(f"partial years excluded from {imports_path.name}: {partial_years}")
    print(f"wrote : {annual_path} ({len(records)} rows)")
    print(f"wrote : {imports_path} ({len(eq)} rows)")
    print("\ndried-equivalent totals by year (tonnes):")
    totals = {}
    for (year, _), v in eq.items():
        totals[year] = totals.get(year, 0.0) + v
    for year in sorted(totals):
        print(f"  {year}: {totals[year] / 1000:,.1f}")


if __name__ == "__main__":
    main()
